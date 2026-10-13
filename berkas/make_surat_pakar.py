"""
Konsep surat dukungan pribadi dari pakar UNS, mengikuti format surat dukungan
pribadi Prof. Nuryani untuk Tim ASTA.

  Nova-Band   : Prof. Nuryani (Instrumentasi Medis), Prof. Pringgo Widyo Laksono
                -> berkas/dukungan
  NUTRICENSE  : Prof. Pringgo Widyo Laksono -> Desktop/NUTRILENSE/surat (bukan repo ini)
  Teks di [[...]] ditandai kuning = mohon dikonfirmasi.

Semuanya KONSEP tanpa tanda tangan: dikirim ke beliau untuk ditinjau, diubah bila
perlu, lalu ditandatangani sendiri. Tanda tangan dari surat ASTA tidak dipakai.

    python berkas/make_surat_pakar.py [NUTRICENSE|Nova-Band]   -> Konsep Surat Dukungan Pribadi <proyek> - <pakar>.docx + .pdf
"""
import importlib.util
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BROWN = RGBColor(0x9a, 0x34, 0x12)
TANGGAL = "Surakarta, ........ Oktober 2026"

PROYEK = {
    "Nova-Band": dict(
        out=Path(__file__).resolve().parent / "dukungan",
        judul="Nova-Band: An AIoT-Driven Upper-Arm Wearable for Real-Time Physiological "
              "Monitoring and Running Performance",
        ketua="Maheswara Nala Wisadewa",
        anggota="Almira Edgina Nareswari Haryadi, Azra Daniswara Tyandaru, Dela Astiara, "
                "Rizkyta Aufa Mutiagassania, Satriya Javas Wibisono",
        pembimbing="Andreas Wegiq Adia Hendix, S.Pd., M.T.",
        sekolah="SMA Negeri 1 Surakarta"),
    "NUTRICENSE": dict(
        out=Path(r"C:\Users\mosto\Desktop\NUTRILENSE\surat"),
        judul="NUTRICENSE: An AI-Powered Smart Food Safety and Nutritional Assessment Box "
              "for School Feeding Programs",
        ketua="[[(nama ketua tim, mohon diisi)]]",
        anggota="Azzahra Arsya Naura Yusron, Jihan Ufaira Nugrahani, Messiamela Atsirhaq, "
                "Muhammad Hayfa Raimondo, Shazia Ula Naira, William Feliciano Charin "
                "[[(keluarkan nama ketua dari daftar ini)]]",
        pembimbing="[[Andreas Wegiq Adia Hendix, S.Pd., M.T.]]",
        sekolah="SMA Negeri 4 Surakarta"),
}
NURYANI = ("Prof. Nuryani, S.Si., M.Si., Ph.D", "Instrumentasi Medis", "Prof. Nuryani")
PRINGGO = ("[[Prof. Dr. Pringgo Widyo Laksono, S.T., M.Eng.]]", "[[Teknik Industri]]", "Prof. Pringgo Widyo Laksono")

# proyek, (nama, bidang, nama pendek), isi paragraf
PAKAR = [
    ("Nova-Band", NURYANI,
     "Saya memberikan masukan dan arahan terkait pengembangan Nova-Band, baik dari aspek konsep dasar "
     "maupun aspek teknis instrumentasi sensor fisiologis (photoplethysmography/PPG dan sensor gerak IMU), "
     "untuk mendukung keandalan pengukuran serta pengembangan alat secara optimal. Pengembangan teknologi "
     "pemantauan fisiologis yang terjangkau seperti Nova-Band memiliki potensi penting dalam membantu pelari "
     "dan komunitas olahraga berlatih secara lebih aman dan terukur. Oleh karena itu, saya mendukung penelitian "
     "dan pengembangan Nova-Band ini sebagai upaya pengembangan teknologi kesehatan olahraga yang bermanfaat "
     "bagi masyarakat Indonesia."),
    ("Nova-Band", PRINGGO,
     "Saya memberikan masukan dan arahan terkait pengembangan Nova-Band, baik dari aspek konsep dasar, "
     "kebutuhan pengguna, maupun rancangan perangkat yang nyaman dan aman dikenakan di lengan atas saat "
     "berlari, untuk mendukung peningkatan kinerja dan pengembangan produk secara optimal. Pengembangan "
     "teknologi seperti Nova-Band memiliki potensi penting dalam membantu pelari dan komunitas olahraga "
     "memantau kondisi tubuh serta mengatur intensitas latihan secara lebih aman. Oleh karena itu, saya "
     "mendukung penelitian dan pengembangan Nova-Band ini sebagai upaya inovasi teknologi olahraga yang "
     "bermanfaat bagi masyarakat Indonesia."),
    ("NUTRICENSE", PRINGGO,
     "Saya memberikan masukan dan arahan terkait pengembangan NUTRICENSE, baik dari aspek konsep dasar, "
     "kebutuhan pengguna, maupun rancangan sistem dan alur kerja alat dalam penyelenggaraan makan di sekolah, "
     "untuk mendukung peningkatan kinerja dan pengembangan produk secara optimal. Pengembangan teknologi "
     "seperti NUTRICENSE memiliki potensi penting dalam membantu program makan bergizi di sekolah memastikan "
     "makanan yang disajikan aman dikonsumsi dan memenuhi kebutuhan gizi peserta didik. Oleh karena itu, saya "
     "mendukung penelitian dan pengembangan NUTRICENSE ini sebagai upaya inovasi teknologi keamanan pangan "
     "dan gizi yang bermanfaat bagi masyarakat Indonesia."),
]


def runs(p, text, bold=False, italic=False, size=None):
    """Add text; segments wrapped in [[...]] get a yellow highlight."""
    parts = text.replace("]]", "[[").split("[[")
    for i, s in enumerate(parts):
        if not s:
            continue
        r = p.add_run(s)
        r.bold, r.italic = bold, italic
        if size:
            r.font.size = Pt(size)
        if i % 2:
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    return p


def par(d, text="", align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=8, **kw):
    p = d.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.15
    if text:
        runs(p, text, **kw)
    return p


def shade(cell, fill):
    tc = cell._tc.get_or_add_tcPr()
    el = OxmlElement("w:shd"); el.set(qn("w:val"), "clear"); el.set(qn("w:fill"), fill); tc.append(el)


def kv_table(d, rows, indent_cm=0.6):
    t = d.add_table(rows=0, cols=3)
    t.alignment = WD_TABLE_ALIGNMENT.LEFT
    t.autofit = False
    widths = (Cm(3.6), Cm(0.5), Cm(12.2))
    ind = OxmlElement("w:tblInd"); ind.set(qn("w:w"), str(int(indent_cm * 567))); ind.set(qn("w:type"), "dxa")
    t._tbl.tblPr.append(ind)
    for k, v, it in rows:
        cells = t.add_row().cells
        for c, w in zip(cells, widths):
            c.width = w
            c.paragraphs[0].paragraph_format.space_after = Pt(2)
        runs(cells[0].paragraphs[0], k)
        runs(cells[1].paragraphs[0], ":")
        runs(cells[2].paragraphs[0], v, italic=it)
    return t


def build(proyek, pakar, isi):
    P = PROYEK[proyek]
    nama, bidang, ttd = pakar
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.59), Cm(27.94)
    s.top_margin, s.bottom_margin = Cm(1.8), Cm(2.0)
    s.left_margin = s.right_margin = Cm(2.3)
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    b = d.add_table(rows=1, cols=1); c = b.rows[0].cells[0]; shade(c, "FFF4E5")
    r = c.paragraphs[0].add_run("KONSEP untuk ditinjau %s. " % ttd)
    r.bold = True; r.font.size = Pt(9); r.font.color.rgb = BROWN
    r = c.paragraphs[0].add_run("Disiapkan oleh Tim Peneliti %s sebagai bahan; isi dapat diubah seluruhnya. "
                                "Bagian berwarna kuning mohon dikonfirmasi. Surat berlaku setelah ditandatangani "
                                "sendiri oleh yang bersangkutan. Hapus kotak ini sebelum ditandatangani." % proyek)
    r.font.size = Pt(9); r.font.color.rgb = BROWN
    par(d, after=4)

    p = par(d, align=WD_ALIGN_PARAGRAPH.CENTER, after=14)
    r = p.add_run("SURAT DUKUNGAN"); r.bold = True; r.underline = True; r.font.size = Pt(13)

    par(d, "Yang bertanda tangan di bawah ini:", after=2)
    kv_table(d, [("Nama", nama, False), ("Bidang keahlian", bidang, False)])
    par(d, after=0)
    par(d, "setelah bertemu dan berdiskusi dengan Tim Peneliti %s, dengan ini secara "
           "pribadi menyatakan dukungan terhadap penelitian dan pengembangan:" % P["sekolah"], after=2)
    kv_table(d, [("Judul", P["judul"], True), ("Ketua Tim", P["ketua"], False), ("Anggota", P["anggota"], False),
                 ("Guru Pembimbing", P["pembimbing"], False), ("Asal Sekolah", P["sekolah"], False)])
    par(d, after=0)
    par(d, isi)
    par(d, "Demikian surat dukungan ini saya buat dengan sebenarnya untuk dapat dipergunakan sebagaimana mestinya.",
        after=18)

    for line, after in ((TANGGAL, 0), ("Yang menyatakan,", 60), (nama, 0)):
        p = par(d, align=WD_ALIGN_PARAGRAPH.LEFT, after=after)
        p.paragraph_format.left_indent = Cm(8.6)
        runs(p, line)

    P["out"].mkdir(parents=True, exist_ok=True)
    f = P["out"] / ("Konsep Surat Dukungan Pribadi %s - %s.docx" % (proyek, ttd))
    d.save(f)
    return f


if __name__ == "__main__":
    only = sys.argv[1:]                               # e.g. NUTRICENSE: build just that project
    made = [build(*p) for p in PAKAR if not only or p[0] in only]
    spec = importlib.util.spec_from_file_location("dispen", r"C:\Users\mosto\Desktop\meguru\tools\make_dispen.py")
    D = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(D)
    D.ke_pdf(made)
    for m in made:
        print("ok", m, "+ .pdf")
