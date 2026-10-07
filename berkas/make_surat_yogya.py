"""
Surat Mersif Academy untuk tim Nova-Band (Yogyakarta):
  1. Dispensasi Maheswara Nala Wisadewa -> Bapak/Ibu Guru SMA Negeri 5 Yogyakarta
  2. Surat pengantar permohonan surat dukungan -> SMAN 5 Yogyakarta, SMAN 1 Ngaglik,
     SMAN 2 Ngaglik, Universitas Negeri Yogyakarta

Memakai templat, kop, dan tanda tangan mentor dari Desktop/meguru (tools/make_dispen.py).
    python berkas/make_surat_yogya.py      -> berkas/mersif/*.docx + *.pdf  (tidak di-commit)
"""
import copy
import importlib.util
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH

MEGURU = Path(r"C:\Users\mosto\Desktop\meguru\tools\make_dispen.py")
spec = importlib.util.spec_from_file_location("dispen", MEGURU)
D = importlib.util.module_from_spec(spec)
spec.loader.exec_module(D)

OUT = Path(__file__).resolve().parent / "mersif"
TANGGAL = "08 Oktober 2026"
HARI = "Kamis, 08 Oktober 2026"
JUDUL = ("“Nova-Band: An AIoT-Driven Upper-Arm Wearable for Real-Time Physiological Monitoring "
         "and Running Performance”")
KONTEKS = ("persiapan mengikuti kompetisi ISIF 2026 yang akan diselenggarakan di Universitas Indonesia, "
           "Depok, dengan judul penelitian " + JUDUL + " yang didampingi oleh Mersif Academy")
TIM = [("Maheswara Nala Wisadewa", "Ketua Tim"), ("Almira Edgina Nareswari Haryadi", "Anggota"),
       ("Azra Daniswara Tyandaru", "Anggota"), ("Dela Astiara", "Anggota"),
       ("Rizkyta Aufa Mutiagassania", "Anggota"), ("Satriya Javas Wibisono", "Anggota")]
DUKUNGAN = [  # nomor, kepada, nama instansi di kalimat, berkas
    ("08.0013/SI/2026", "Yth. Kepala SMA Negeri 5 Yogyakarta", "SMA Negeri 5 Yogyakarta", "SMAN5_Yogyakarta"),
    ("08.0014/SI/2026", "Yth. Kepala SMA Negeri 1 Ngaglik", "SMA Negeri 1 Ngaglik", "SMAN1_Ngaglik"),
    ("08.0015/SI/2026", "Yth. Kepala SMA Negeri 2 Ngaglik", "SMA Negeri 2 Ngaglik", "SMAN2_Ngaglik"),
    ("08.0016/SI/2026", "Yth. Rektor Universitas Negeri Yogyakarta", "Universitas Negeri Yogyakarta", "UNY"),
]


def surat(nomor, perihal, kepada, pembuka, isi, tabel, kolom3, nama_berkas, lampiran="-"):
    doc = Document(D.TEMPLAT)
    kop = doc.tables[0]
    D.tulis(kop.cell(0, 1).paragraphs[0], f": {nomor}")
    D.tulis(kop.cell(0, 2).paragraphs[0], TANGGAL)
    D.tulis(kop.cell(1, 1).paragraphs[0], f": {lampiran}")
    D.tulis(kop.cell(2, 1).paragraphs[0], f": {perihal}")
    D.tulis(kop.cell(3, 1).paragraphs[0], f": {kepada}")
    D.tulis(D.cari(doc, "Dalam rangka"), pembuka)
    D.tulis(D.cari(doc, "Bermaksud"), isi)
    D.tulis(D.cari(doc, "Atas perhatian"), "Atas perhatian dan dukungan Bapak/Ibu kami ucapkan terima kasih.")
    D.tulis(D.cari(doc, "Senin, 01 September"), HARI)
    tb = doc.tables[1]
    D.tulis(tb.cell(0, 2).paragraphs[0], kolom3)
    contoh = copy.deepcopy(tb.rows[1]._tr)
    for row in list(tb.rows)[1:]:
        tb._tbl.remove(row._tr)
    for i, (a, b) in enumerate(tabel, 1):
        tb._tbl.append(copy.deepcopy(contoh))
        for c, v in zip(tb.rows[-1].cells, (f"{i}.", f" {a}", b)):
            D.tulis(c.paragraphs[0], v)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    paras = doc.paragraphs
    akhir = max(i for i, p in enumerate(paras) if p.text.strip())
    for p in paras[akhir + 1:]:
        if not p.text.strip() and not p._p.xpath(".//w:drawing|.//w:sectPr"):
            p._p.getparent().remove(p._p)
    OUT.mkdir(exist_ok=True)
    dst = OUT / f"{nama_berkas}.docx"
    doc.save(dst)
    return dst


hasil = []
# 1. dispensasi Maheswara
hasil.append(surat(
    "08.0012/SI/2026", "Permohonan Izin Berkegiatan Persiapan Lomba ISIF 2026",
    "Yth. Bapak/Ibu Guru SMA Negeri 5 Yogyakarta",
    f"Dalam rangka {KONTEKS}, siswa kami dengan informasi sebagai berikut :",
    f"Bermaksud untuk mengikuti kegiatan persiapan lomba ISIF 2026 yang akan dilaksanakan pada hari {HARI}. "
    "Maka dari itu kami selaku mentor dan pimpinan lembaga Mersif Academy dengan surat ini meminta dukungan "
    "kepada Bapak/Ibu Guru untuk mengizinkan siswa kami melaksanakan kegiatan tersebut pada hari "
    f"{HARI}. Demikian kami sampaikan, besar harapan kami Bapak/Ibu Guru dapat mendukung kegiatan siswa "
    "yang bersangkutan agar mendapat hasil dan prestasi yang maksimal dalam lomba ISIF 2026.",
    [("Maheswara Nala Wisadewa", "SMA Negeri 5 Yogyakarta")], "Sekolah",
    "Dispensasi_NovaBand_Maheswara_SMAN5Yogyakarta_2026-10-08"))

# 2. surat pengantar permohonan dukungan
for nomor, kepada, inst, kode in DUKUNGAN:
    hasil.append(surat(
        nomor, "Permohonan Surat Dukungan Penelitian Tim Nova-Band", kepada,
        f"Dalam rangka {KONTEKS}, tim kami dengan susunan sebagai berikut :",
        f"Bermaksud memohon dukungan dari {inst} berupa surat dukungan (letter of support) bagi penelitian "
        "Nova-Band, wearable lengan atas untuk memantau detak jantung dan performa lari secara real-time dengan "
        "peringatan dini. Surat tersebut akan kami lampirkan pada lomba ISIF 2026, dan tim kami bersedia "
        f"mendemonstrasikan purwarupanya bila diperlukan. Besar harapan kami {inst} berkenan memberikan "
        "dukungan agar tim Nova-Band meraih hasil yang maksimal.",
        TIM, "Peran", f"Pengantar_Dukungan_NovaBand_{kode}_2026-10-08",
        lampiran="1 (satu) berkas konsep surat dukungan"))

D.ke_pdf(hasil)
for p in hasil:
    print(f"  {p.name}  +  .pdf")
