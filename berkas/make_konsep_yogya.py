"""
Konsep surat dukungan Nova-Band untuk 4 instansi di Yogyakarta, lampiran surat pengantar
Mersif Academy (berkas/make_surat_yogya.py). Pola sama dengan konsep Dinkes/Dispora
(berkas/make_surat_novaband.py): spanduk KONSEP, tempat kop instansi, bagian bertitik,
tanda air "KONSEP" di PDF. Instansi menyalin ke kop resminya, memberi nomor, ttd & cap.

    python berkas/make_konsep_yogya.py   -> berkas/mersif/Konsep_Dukungan_*.docx + .pdf
"""
import importlib.util
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location("nb", HERE / "make_surat_novaband.py")
NB = importlib.util.module_from_spec(spec)
spec.loader.exec_module(NB)
spec = importlib.util.spec_from_file_location("dispen", r"C:\Users\mosto\Desktop\meguru\tools\make_dispen.py")
DP = importlib.util.module_from_spec(spec)
spec.loader.exec_module(DP)

DOTS, GREY, BROWN = NB.DOTS, NB.GREY, NB.BROWN
JUDUL = NB.JUDUL
TIM = NB.TIM
OUT = HERE / "mersif"
PERIHAL_PENGANTAR = "Permohonan Surat Dukungan Penelitian Tim Nova-Band"

BATASAN = ("Dukungan ini diberikan dalam rangka pembinaan kegiatan penelitian dan inovasi pelajar, dan tidak "
           "merupakan pernyataan kelayakan klinis, rekomendasi komersial atas produk, maupun izin edar alat "
           "kesehatan. Setiap uji coba yang melibatkan siswa atau peserta wajib dilaksanakan dengan persetujuan "
           "peserta (dan orang tua/wali bagi peserta di bawah umur), mengutamakan keselamatan, serta menjaga "
           "kerahasiaan data pribadi.")

SEKOLAH_DUKUNGAN = lambda nama: [
    ("Dukungan moril", ", berupa apresiasi dan dorongan kepada tim peneliti untuk melanjutkan penelitian dan "
                       "pengembangan Nova-Band secara bertanggung jawab hingga kompetisi ISIF 2026."),
    ("Dukungan fasilitasi", ", berupa kesempatan pengenalan dan uji coba terbatas Nova-Band di lingkungan %s, "
                            "misalnya pada kegiatan Pendidikan Jasmani, Olahraga, dan Kesehatan (PJOK) atau "
                            "ekstrakurikuler olahraga, sesuai jadwal dan ketentuan sekolah." % nama),
    ("Dukungan masukan", ", berupa kesempatan berdiskusi dengan guru PJOK dan pembina olahraga sekolah untuk "
                         "memberi masukan mengenai kebutuhan pelari pelajar serta aspek keselamatan berolahraga."),
]

INSTANSI = {
    "SMAN5_Yogyakarta": dict(
        nama="SMA Negeri 5 Yogyakarta", tempat="Yogyakarta", nomor="08.0013/SI/2026",
        pandangan=("SMA Negeri 5 Yogyakarta memandang bahwa penelitian Nova-Band, yang diketuai oleh siswa kami, "
                   "Maheswara Nala Wisadewa, merupakan wujud semangat inovasi pelajar di bidang teknologi olahraga "
                   "dan kesehatan yang sejalan dengan upaya sekolah menumbuhkan budaya riset serta gaya hidup aktif "
                   "dan aman. Dukungan yang kami berikan meliputi:"),
        dukungan=SEKOLAH_DUKUNGAN("SMA Negeri 5 Yogyakarta"),
        jabatan=["Kepala SMA Negeri 5 Yogyakarta,"], pejabat="( nama Kepala Sekolah )", nip=True),
    "SMAN1_Ngaglik": dict(
        nama="SMA Negeri 1 Ngaglik", tempat="Sleman", nomor="08.0014/SI/2026",
        pandangan=("SMA Negeri 1 Ngaglik memandang bahwa pengembangan teknologi pemantauan kondisi tubuh yang "
                   "terjangkau untuk membantu pelajar berolahraga secara aman dan terukur sejalan dengan pembinaan "
                   "pendidikan jasmani, kesehatan, dan budaya inovasi di kalangan pelajar. Dukungan yang kami "
                   "berikan meliputi:"),
        dukungan=SEKOLAH_DUKUNGAN("SMA Negeri 1 Ngaglik"),
        jabatan=["Kepala SMA Negeri 1 Ngaglik,"], pejabat="( nama Kepala Sekolah )", nip=True),
    "SMAN2_Ngaglik": dict(
        nama="SMA Negeri 2 Ngaglik", tempat="Sleman", nomor="08.0015/SI/2026",
        pandangan=("SMA Negeri 2 Ngaglik memandang bahwa pengembangan teknologi pemantauan kondisi tubuh yang "
                   "terjangkau untuk membantu pelajar berolahraga secara aman dan terukur sejalan dengan pembinaan "
                   "pendidikan jasmani, kesehatan, dan budaya inovasi di kalangan pelajar. Dukungan yang kami "
                   "berikan meliputi:"),
        dukungan=SEKOLAH_DUKUNGAN("SMA Negeri 2 Ngaglik"),
        jabatan=["Kepala SMA Negeri 2 Ngaglik,"], pejabat="( nama Kepala Sekolah )", nip=True),
    "UNY": dict(
        nama="Universitas Negeri Yogyakarta", tempat="Yogyakarta", nomor="08.0016/SI/2026",
        pandangan=("Universitas Negeri Yogyakarta memandang bahwa penelitian Nova-Band relevan dengan "
                   "pengembangan ilmu keolahragaan dan kesehatan, khususnya pemanfaatan teknologi untuk pemantauan "
                   "beban latihan dan keselamatan pelari, serta sejalan dengan komitmen universitas dalam membina "
                   "talenta muda di bidang riset dan inovasi. Dukungan yang kami berikan meliputi:"),
        dukungan=[
            ("Dukungan moril", ", berupa apresiasi dan dorongan kepada tim peneliti pelajar untuk melanjutkan "
                               "penelitian dan pengembangan Nova-Band secara ilmiah dan bertanggung jawab."),
            ("Dukungan kepakaran", ", berupa kesempatan diskusi dan pemberian masukan dari dosen/peneliti di bidang "
                                   "ilmu keolahragaan, fisiologi olahraga, dan teknologi terkait metodologi "
                                   "pengujian serta interpretasi data Nova-Band."),
            ("Dukungan fasilitasi", ", berupa kemungkinan pemanfaatan fasilitas laboratorium atau lapangan untuk "
                                    "pengujian pembanding, sesuai jadwal, prosedur, dan ketentuan yang berlaku di "
                                    "Universitas Negeri Yogyakarta."),
        ],
        jabatan=["Rektor Universitas Negeri Yogyakarta", "atau pejabat yang ditunjuk,"],
        pejabat="( nama pejabat )", nip=True),
}


def build(key):
    I = INSTANSI[key]
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21.0), Cm(29.7)
    s.top_margin = Cm(1.5); s.bottom_margin = Cm(1.4); s.left_margin = Cm(2.3); s.right_margin = Cm(2.3)
    st = d.styles["Normal"]; st.font.name = "Times New Roman"; st.font.size = Pt(11)

    b = d.add_table(rows=1, cols=1); c = b.rows[0].cells[0]; NB.shade(c, "FFF4E5")
    r = c.paragraphs[0].add_run("KONSEP untuk ditinjau %s. " % I["nama"])
    r.bold = True; r.font.size = Pt(8.5); r.font.color.rgb = BROWN
    r2 = c.paragraphs[0].add_run("Disiapkan oleh Tim Nova-Band (Mersif Academy) sebagai bahan; isi dapat diubah "
                                 "seluruhnya. Surat berlaku setelah diterbitkan di kop surat resmi, diberi nomor, "
                                 "ditandatangani, dan dicap. Hapus kotak ini dan catatan di akhir saat diterbitkan.")
    r2.font.size = Pt(8.5); r2.font.color.rgb = BROWN

    k = d.add_table(rows=1, cols=1); kc = k.rows[0].cells[0]; NB.shade(kc, "F3F4F6")
    kp = kc.paragraphs[0]; kp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kr = kp.add_run("[ KOP SURAT RESMI %s ]" % I["nama"].upper()); kr.bold = True; kr.font.size = Pt(10)
    kr.font.color.rgb = GREY
    g = d.add_paragraph(); pPr = g._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    bt.set(qn("w:val"), "double"); bt.set(qn("w:sz"), "18"); bt.set(qn("w:space"), "1"); bt.set(qn("w:color"), "000000")
    bdr.append(bt); pPr.append(bdr); g.paragraph_format.space_after = Pt(8)

    meta = d.add_table(rows=3, cols=3)
    for i, (a, v, bold) in enumerate([("Nomor", ".................. / .................. / 2026", False),
                                       ("Lampiran", "\u2013", False),
                                       ("Perihal", "Surat Dukungan Penelitian Tim Nova-Band", True)]):
        meta.rows[i].cells[0].paragraphs[0].add_run(a)
        meta.rows[i].cells[1].paragraphs[0].add_run(":")
        rr = meta.rows[i].cells[2].paragraphs[0].add_run(v); rr.bold = bold
    NB.fix_table(meta, (2.4, 0.5, 13.5))

    NB.par(d, I["tempat"] + ", " + DOTS + " 2026", align=WD_ALIGN_PARAGRAPH.RIGHT, after=8)
    for i, (t, bold) in enumerate([("Yth.", False), ("Pimpinan Mersif Academy", True),
                                   ("c.q. Tim Nova-Band", False), ("di tempat", False)]):
        NB.par(d, t, bold=bold, align=WD_ALIGN_PARAGRAPH.LEFT, after=0 if i < 3 else 8)

    NB.par(d, "Dengan hormat,")
    NB.par(d, "Menindaklanjuti surat Mersif Academy Nomor %s tanggal 08 Oktober 2026 perihal %s, dengan ini "
              "%s menyampaikan dukungan terhadap kegiatan penelitian dan pengembangan berikut:"
           % (I["nomor"], PERIHAL_PENGANTAR, I["nama"]))
    rows = [("Judul", JUDUL), ("Ketua Tim", TIM[0]), ("Anggota", ", ".join(TIM[1:])),
            ("Pendamping", "Mersif Academy"), ("Kompetisi", "ISIF 2026, Universitas Indonesia, Depok")]
    t = d.add_table(rows=len(rows), cols=3)
    for i, (a, v) in enumerate(rows):
        t.rows[i].cells[0].paragraphs[0].add_run(a)
        t.rows[i].cells[1].paragraphs[0].add_run(":")
        rv = t.rows[i].cells[2].paragraphs[0].add_run(v); rv.italic = (a == "Judul")
    NB.fix_table(t, (3.2, 0.5, 12.7))
    d.add_paragraph().paragraph_format.space_after = Pt(0)

    NB.par(d, I["pandangan"])
    for i, (lead, rest) in enumerate(I["dukungan"], 1):
        p = d.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.left_indent = Cm(0.8); p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(3)
        p.add_run("%d. " % i); rb = p.add_run(lead); rb.bold = True; p.add_run(rest)
    NB.par(d, BATASAN)
    NB.par(d, "Demikian surat dukungan ini dibuat untuk dapat dipergunakan sebagaimana mestinya.", after=8)

    sg = d.add_table(rows=1, cols=2)
    cell = sg.rows[0].cells[1]
    cell.paragraphs[0].add_run(I["jabatan"][0])
    for extra in I["jabatan"][1:]:
        cell.add_paragraph(extra)
    pr = cell.add_paragraph(); rr = pr.add_run("[tanda tangan & cap \u2014 dibubuhkan oleh instansi]")
    rr.italic = True; rr.font.size = Pt(8.5); rr.font.color.rgb = GREY
    cell.add_paragraph()
    pn = cell.add_paragraph(); rn = pn.add_run(I["pejabat"]); rn.bold = True; rn.underline = True
    if I["nip"]:
        cell.add_paragraph("NIP " + DOTS)

    n = d.add_table(rows=1, cols=1); nc = n.rows[0].cells[0]; NB.shade(nc, "F6F8FA")
    nr = nc.paragraphs[0].add_run("Catatan. "); nr.bold = True; nr.font.size = Pt(8)
    nr2 = nc.paragraphs[0].add_run(
        "Bagian bertitik (nomor, tanggal, nama dan NIP pejabat) sengaja dikosongkan. Rumusan dukungan, terutama "
        "butir fasilitasi dan kalimat batasan, dapat disesuaikan dengan kebijakan dan kewenangan instansi. "
        "Catatan ini, spanduk konsep, dan tanda air dihapus saat diterbitkan.")
    nr2.font.size = Pt(8); nr2.font.color.rgb = GREY

    # compact: single line spacing everywhere, no space under table-cell paragraphs
    for p_ in d.paragraphs:
        p_.paragraph_format.line_spacing = 1.0
        if p_.paragraph_format.space_after and p_.paragraph_format.space_after > Pt(4):
            p_.paragraph_format.space_after = Pt(4)
    for tb in d.tables:
        for row in tb.rows:
            for cl in row.cells:
                for p_ in cl.paragraphs:
                    p_.paragraph_format.space_after = Pt(0)
                    p_.paragraph_format.space_before = Pt(0)
                    p_.paragraph_format.line_spacing = 1.0
    OUT.mkdir(exist_ok=True)
    out = OUT / ("Konsep_Dukungan_NovaBand_%s.docx" % key)
    d.save(out)
    return out


def watermark(pdf):
    import fitz
    doc = fitz.open(pdf)
    for page in doc:
        r = page.rect
        page.insert_text((r.width * 0.2, r.height * 0.62), "KONSEP", fontsize=110, color=(0.93, 0.91, 0.90),
                         rotate=0, morph=(fitz.Point(r.width / 2, r.height / 2), fitz.Matrix(-30)), overlay=False)
    tmp = str(pdf) + ".tmp"
    doc.save(tmp)
    doc.close()
    os.replace(tmp, pdf)


if __name__ == "__main__":
    made = [build(k) for k in INSTANSI]
    DP.ke_pdf(made)
    for m in made:
        watermark(m.with_suffix(".pdf"))
        print("ok", m.name)
