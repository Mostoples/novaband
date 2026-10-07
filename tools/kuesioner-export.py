"""
Download the runner questionnaire (kuesioner.html -> Firestore kuesioner_pelari)
and build an Excel workbook: raw answers, a per-question summary, and the
Nova-Band feature ranking (Part B, mean importance 1-5).

    python tools/kuesioner-export.py            # -> data/kuesioner/kuesioner_pelari_<date>.xlsx (+ .csv)

Reading needs admin credentials (browsers are blocked by firestore.rules):
the gcloud login, or the Firebase CLI login (`firebase login`). Rows whose
name starts with "UJI COBA" or whose suggestion starts with "[UJI]" are test
submissions and are skipped (the admin page skips "[UJI]" too).
The output holds respondents' answers: it stays out of git (data/ is ignored).
"""
import csv
import datetime as dt
import json
import os
import shutil
import subprocess
import urllib.request
from collections import Counter

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "data", "kuesioner")
PROJECT = "novaband-id"
URL = ("https://firestore.googleapis.com/v1/projects/%s/databases/(default)/documents/kuesioner_pelari"
       "?pageSize=300" % PROJECT)

COLS = ["dibuat", "nama", "usia", "gender", "kota", "kategori",
        "a1_frekuensi", "a2_jarak", "a3_perangkat", "a4_penting_hr", "a5_keluhan", "a6_akurasi",
        "a7_kendala", "a8_tahu_batas", "a9_pendamping", "a10_harga",
        "b1_hr_akurat", "b2_peringatan", "b3_notif_pelatih", "b4_readiness", "b5_cedera",
        "b6_data_lari", "b7_layar", "b8_ai_coach", "b9_komunitas", "b10_nyaman",
        "saran", "durasi_detik"]
FITUR = {
    "b1_hr_akurat": "Detak jantung real-time akurat (lengan atas)",
    "b2_peringatan": "Peringatan dini saat HR melewati batas",
    "b3_notif_pelatih": "Notifikasi ke pelatih / keluarga / grup",
    "b4_readiness": "Readiness Score harian",
    "b5_cedera": "Analisis beban otot & risiko cedera (IMU)",
    "b6_data_lari": "Data lari lengkap (pace, jarak, kadensi, kalori)",
    "b7_layar": "Layar sentuh di band",
    "b8_ai_coach": "AI coach Nova (rekomendasi latihan)",
    "b9_komunitas": "Fitur komunitas / peta kelompok pelari",
    "b10_nyaman": "Nyaman, ringan, baterai seharian",
}


def token():
    g = shutil.which("gcloud") or shutil.which("gcloud.cmd")
    if g:
        r = subprocess.run([g, "auth", "print-access-token"], capture_output=True, text=True, shell=os.name == "nt")
        if r.returncode == 0 and r.stdout.strip():
            return r.stdout.strip()
    # Firebase CLI: any command refreshes the stored access token
    subprocess.run("firebase projects:list", shell=True, capture_output=True)
    cfg = json.load(open(os.path.expanduser("~/.config/configstore/firebase-tools.json")))
    return cfg["tokens"]["access_token"]


def val(v):
    if "stringValue" in v:
        return v["stringValue"]
    if "integerValue" in v:
        return int(v["integerValue"])
    if "booleanValue" in v:
        return v["booleanValue"]
    if "timestampValue" in v:
        return v["timestampValue"]
    if "arrayValue" in v:
        return [val(x) for x in v["arrayValue"].get("values", [])]
    return None


def fetch():
    tok, rows, page = token(), [], ""
    while True:
        req = urllib.request.Request(URL + (("&pageToken=" + page) if page else ""),
                                     headers={"Authorization": "Bearer " + tok})
        data = json.load(urllib.request.urlopen(req))
        for d in data.get("documents", []):
            rows.append({k: val(v) for k, v in d.get("fields", {}).items()})
        page = data.get("nextPageToken")
        if not page:
            return rows


def main():
    rows = [r for r in fetch() if not str(r.get("nama", "")).upper().startswith("UJI COBA")
            and not str(r.get("saran", "")).startswith("[UJI]")]
    rows.sort(key=lambda r: r.get("dibuat", ""))
    os.makedirs(OUT, exist_ok=True)
    stamp = dt.date.today().isoformat()
    flat = lambda v: "; ".join(v) if isinstance(v, list) else ("" if v is None else v)

    path_csv = os.path.join(OUT, "kuesioner_pelari_%s.csv" % stamp)
    with open(path_csv, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(COLS)
        for r in rows:
            w.writerow([flat(r.get(c)) for c in COLS])

    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = "Jawaban"
    ws.append(COLS)
    for r in rows:
        ws.append([flat(r.get(c)) for c in COLS])
    head = PatternFill("solid", fgColor="7B2021")
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    ws.freeze_panes = "B2"

    # Part B ranking
    rk = wb.create_sheet("Prioritas Fitur")
    rk.append(["Peringkat", "Fitur Nova-Band", "Rata-rata (1-5)", "% penting (4-5)", "n"])
    stats = []
    for k, label in FITUR.items():
        xs = [r[k] for r in rows if isinstance(r.get(k), int)]
        if xs:
            stats.append((sum(xs) / len(xs), sum(x >= 4 for x in xs) / len(xs), len(xs), label))
    for i, (m, p, n, label) in enumerate(sorted(stats, reverse=True), 1):
        rk.append([i, label, round(m, 2), "%.0f%%" % (p * 100), n])
    for c in rk[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head

    # per-question counts
    sm = wb.create_sheet("Rekap")
    sm.append(["Pertanyaan", "Jawaban", "Jumlah", "%"])
    for c in sm[1]:
        c.font, c.fill = Font(bold=True, color="FFFFFF"), head
    for col in COLS[2:-2]:
        cnt = Counter()
        for r in rows:
            v = r.get(col)
            for x in (v if isinstance(v, list) else [v]):
                if x not in (None, ""):
                    cnt[x] += 1
        sm.append([col, "", "", ""])
        sm.cell(sm.max_row, 1).font = Font(bold=True)
        for ans, k in sorted(cnt.items(), key=lambda kv: -kv[1]):
            sm.append(["", ans, k, "%.0f%%" % (100 * k / max(len(rows), 1))])
    for s in (ws, rk, sm):
        for col in s.columns:
            s.column_dimensions[col[0].column_letter].width = min(48, max(10, max(len(str(c.value or "")) for c in col) + 2))
    path_xlsx = os.path.join(OUT, "kuesioner_pelari_%s.xlsx" % stamp)
    wb.save(path_xlsx)
    print("%d respons" % len(rows))
    for m, p, n, label in sorted(stats, reverse=True)[:5]:
        print("  %.2f  %3.0f%%  %s" % (m, p * 100, label))
    print(path_xlsx)
    print(path_csv)


if __name__ == "__main__":
    main()
